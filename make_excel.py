"""Build results/SECOND_SCD_Results.xlsx: every result of this project in one clean workbook.

Run make_docs.py first (it draws the charts in docs/figures/), then:
    .venv/bin/python make_excel.py

Design rules, so nothing overlaps in Excel, Numbers, Google Sheets or file previews:
- every sheet has the same layout: title, "what this sheet shows", a highlighted key takeaway,
  then tables, notes and charts stacked top to bottom with measured spacing;
- charts are the PNG figures from docs/figures (they look the same in every program);
- raw numbers come from results/*.json; derived numbers (gains, means, std, ranks, win counts,
  error rates) are Excel formulas, and their computed values are also stored in the file.
"""
import glob
import json
import math
import os

import numpy as np
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from PIL import Image as PILImage

ROOT = os.path.dirname(os.path.abspath(__file__))
R = os.path.join(ROOT, "results")
FIG = os.path.join(ROOT, "docs", "figures")
OUT = os.path.join(R, "SECOND_SCD_Results.xlsx")
CLASSES = ["no-change", "non-veg ground", "tree", "low vegetation", "water", "building", "playground"]
SEM = CLASSES[1:]

# ============================================================================ data
runs = {os.path.basename(p)[:-5]: json.load(open(p))
        for p in sorted(glob.glob(os.path.join(R, "[AB]_*.json")))}
imb = json.load(open(os.path.join(R, "imbalance_stats.json")))
best_arch = open(os.path.join(R, "best_arch.txt")).read().split()[-1]

ARCH = {"early_fusion": ("Early Fusion", "HRSCD strategy 2 (Daudt et al., 2019)", 11.36,
                         "Puts both images together as one 6-channel input; one network predicts everything"),
        "sscd": ("SSCD", "SSCD-l (Ding et al., 2022)", 11.58,
                 "Siamese: the same encoder reads each image separately; separate change branch"),
        "bisrnet": ("Bi-SRNet-lite", "Bi-SRNet (Ding et al., 2022)", 11.91,
                    "SSCD + the two images 'look at each other' (attention) + consistency loss")}
TECH = {  # (loss, sampler) -> label, plain-language description, reference, link
    ("ce", "uniform"): ("CE (baseline)", "Standard loss, every pixel counts the same", "Standard", ""),
    ("wce", "uniform"): ("Weighted CE", "Rare classes get a bigger weight in the loss", "Common practice", ""),
    ("median", "uniform"): ("Median-freq", "Class weight = median frequency / class frequency",
                            "Eigen & Fergus, ICCV 2015", "https://arxiv.org/abs/1411.4734"),
    ("cb", "uniform"): ("Class-balanced", "Weights from the 'effective number' of samples per class",
                        "Cui et al., CVPR 2019", "https://arxiv.org/abs/1901.05555"),
    ("focal", "uniform"): ("Focal", "Pays less attention to pixels that are already easy",
                           "Lin et al., ICCV 2017", "https://arxiv.org/abs/1708.02002"),
    ("dice", "uniform"): ("CE + Dice", "Adds a loss that measures overlap, so small classes count more",
                          "Milletari et al., 3DV 2016", "https://arxiv.org/abs/1606.04797"),
    ("ohem", "uniform"): ("OHEM", "Trains only on the hardest 25% of pixels",
                          "Shrivastava et al., CVPR 2016", "https://arxiv.org/abs/1604.03540"),
    ("ce", "rare"): ("CE + rare sampling", "Same loss; images with rare classes are shown more often",
                     "Gupta et al. (LVIS), CVPR 2019", "https://arxiv.org/abs/1908.03195"),
    ("combo", "uniform"): ("WCE + Dice", "Weighted CE and Dice together",
                           "Weighted CE + Milletari et al.", "https://arxiv.org/abs/1606.04797"),
    ("combo", "rare"): ("WCE + Dice + rare sampling", "WCE + Dice, and rare-class images shown more often",
                        "As above + LVIS sampling", "https://arxiv.org/abs/1908.03195"),
}


def info(tag):
    r = runs[tag]
    a, t = r["args"], r["test"]
    return dict(tag=tag, model=a["model"], key=(a["loss"], a["sampler"]),
                label=TECH[(a["loss"], a["sampler"])][0], t=t, pc=t["per_class_IoU"],
                val=max(h["SeK"] for h in r["history"]), ep=r["best_epoch"],
                mins=r["train_minutes"], hist=r["history"], conf=np.array(r["confusion"]))


A = sorted([info(t) for t in runs if t.startswith("A_")],
           key=lambda x: ["early_fusion", "sscd", "bisrnet"].index(x["model"]))
base = next(x for x in A if x["model"] == best_arch)
ef = next(x for x in A if x["model"] == "early_fusion")
B = sorted([info(t) for t in runs if t.startswith("B_")], key=lambda x: -x["t"]["SeK"])
BB = [base] + B
ALL = A + B
best = max(ALL, key=lambda x: x["t"]["SeK"])
full = lambda x: f"{ARCH[x['model']][0]} + {x['label']}"

CV = {}
for p in sorted(glob.glob(os.path.join(R, "cv", "*_f[0-9].json"))):
    r = json.load(open(p))
    CV.setdefault(r["tag"].rsplit("_f", 1)[0], {})[r["args"]["fold"]] = r
CV_ORDER = ["A_early_fusion_ce", "A_sscd_ce", "A_bisrnet_ce", "B_bisrnet_combo", "B_bisrnet_dice",
            "B_bisrnet_wce", "B_bisrnet_median", "B_bisrnet_cb", "B_bisrnet_focal",
            "B_bisrnet_ce_rare", "B_bisrnet_combo_rare", "B_bisrnet_ohem"]
K = 3
cv_done = sum(len(v) for v in CV.values())
cv_complete = [t for t in CV_ORDER if len(CV.get(t, {})) == K]


def cv_label(tag):
    model = "early_fusion" if "early" in tag else ("sscd" if "sscd" in tag else "bisrnet")
    if tag.startswith("A_"):
        key = ("ce", "uniform")
    else:
        rest = tag.split("_", 2)[2]
        key = (rest.replace("_rare", ""), "rare" if rest.endswith("_rare") else "uniform")
    return f"{ARCH[model][0]} + {TECH[key][0]}"


def cv_stats(tag):
    rs = [CV[tag][k] for k in range(K)]
    be = [r["best_epoch"] - 1 for r in rs]
    va = np.array([[h["SeK"] for h in r["history"]] for r in rs])
    tr = np.array([[h["train_SeK"] for h in r["history"]] for r in rs])
    vl = np.array([[h["val_loss"] for h in r["history"]] for r in rs])
    rise = float(np.mean((vl[:, -1] - vl.min(1)) / np.abs(vl.min(1))))
    g_end = float((tr[:, -1] - va[:, -1]).mean())
    slope = float(np.mean([np.polyfit(np.arange(5), v[-5:], 1)[0] for v in va]))
    if va.max() < 5:
        diag = "Failed to learn"
    elif rise > 0.10 and g_end > 3:
        diag = "Overfitting"
    elif slope > 0.15 and g_end < 2:
        diag = "Underfitting (still improving)"
    elif g_end > 3 or rise > 0.10:
        diag = "Mild overfitting"
    else:
        diag = "Good fit"
    last_best = all(b == va.shape[1] - 1 for b in be)
    return dict(test=[r["test"]["SeK"] for r in rs], val=float(np.mean([va[i, b] for i, b in enumerate(be)])),
                train=float(np.mean([tr[i, b] for i, b in enumerate(be)])), rise=rise, diag=diag,
                last_best=last_best, rs=rs, be=be)


CVS = {t: cv_stats(t) for t in cv_complete}

LIT = [("HRSCD-str.2", 85.49, 64.43, 10.69, 49.22), ("HRSCD-str.4", 86.62, 71.15, 18.80, 58.21),
       ("SSCD-l", 87.19, 72.60, 21.86, 61.22), ("Bi-SRNet", 87.84, 73.41, 23.22, 62.61),
       ("SCanNet", 87.86, 73.42, 23.94, 63.66), ("ChangeMamba", 88.12, 73.68, 24.11, 64.03)]

# ============================================================================ style & layout helpers
NAVY, INK2 = "0D366B", "52514E"
F = lambda **k: Font(name="Arial", **{"size": 10, **k})
HEAD = PatternFill("solid", fgColor="184F95")
BAND = PatternFill("solid", fgColor="F3F6FB")
BEST = PatternFill("solid", fgColor="DDF1E6")
TAKE = PatternFill("solid", fgColor="E8F3EC")
NOTE = PatternFill("solid", fgColor="FFF8E6")
NONE = PatternFill()
thin = Side(style="thin", color="D0D0D0")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
LEFT_BAR = Border(left=Side(style="thick", color="1BAF7A"))
PT_PER_LINE = 13.5

wb = Workbook()
wb.remove(wb.active)
REFS = {}                # cell addresses used by formulas on other sheets


def width_of(ws, c):
    return ws.column_dimensions[L(c)].width or 8.43


def lines_needed(text, chars):
    text = str(text)
    return sum(max(1, math.ceil(len(part) / max(chars, 1))) for part in text.split("\n"))


def block(ws, row, text, ncols, font, fill=NONE, border=None, min_h=15):
    """A merged, wrapped text block across ncols columns; the row height fits the text."""
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row, 1, text)
    c.font = font
    c.fill = fill
    c.alignment = Alignment(wrap_text=True, vertical="top", indent=1 if fill is not NONE else 0)
    if border:
        c.border = border
    chars = sum(width_of(ws, k) for k in range(1, ncols + 1)) * 1.05 * 10 / font.size - 4
    ws.row_dimensions[row].height = max(min_h, lines_needed(text, chars) * PT_PER_LINE * font.size / 10 + 6)
    return row + 1


def new_sheet(name, title, what, takeaway, widths):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    for i, w in enumerate(widths):
        ws.column_dimensions[L(i + 1)].width = w
    n = len(widths)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    r = block(ws, 1, title, n, F(size=16, bold=True, color=NAVY), min_h=26)
    r = block(ws, r, "What this sheet shows: " + what, n, F(size=10, color=INK2))
    r = block(ws, r, "Key takeaway: " + takeaway, n, F(size=11, bold=True, color="0B4F2F"), fill=TAKE,
              border=LEFT_BAR, min_h=20)
    return ws, r + 1, n


def heading(ws, row, text, n):
    ws.cell(row, 1, text).font = F(size=12, bold=True, color=NAVY)
    ws.row_dimensions[row].height = 20
    return row + 1


def note(ws, row, text, n):
    return block(ws, row, text, n, F(size=9, italic=True, color=INK2), fill=NOTE) + 1


def table(ws, row, headers, rows, fmts=None, best_row=None, wrap_cols=()):
    """Styled table starting in column A; row heights fit wrapped text. Returns (first, last, next_row)."""
    hh = 1
    for j, h in enumerate(headers):
        c = ws.cell(row, 1 + j, h)
        c.font = F(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.border = BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        hh = max(hh, lines_needed(h, width_of(ws, 1 + j) * 0.95 - 1))
    ws.row_dimensions[row].height = hh * PT_PER_LINE + 8
    for i, rv in enumerate(rows):
        rr = row + 1 + i
        nl = 1
        for j, v in enumerate(rv):
            c = ws.cell(rr, 1 + j, v)
            is_text = isinstance(v, str) and not v.startswith("=")
            wrap = j == 0 or j in wrap_cols
            c.font = F(bold=(best_row == i))
            c.border = BOX
            c.fill = BEST if best_row == i else (BAND if i % 2 else NONE)
            c.alignment = Alignment(vertical="center", wrap_text=wrap,
                                    horizontal="left" if is_text or j == 0 else "right")
            if fmts and j < len(fmts) and fmts[j]:
                c.number_format = fmts[j]
            if wrap and is_text:
                nl = max(nl, lines_needed(v, width_of(ws, 1 + j) * 1.05 - 2))
        ws.row_dimensions[rr].height = max(16, nl * PT_PER_LINE + 5)
    return row + 1, row + len(rows), row + len(rows) + 2


def picture(ws, row, fname, n, max_px=1000):
    """Insert a PNG from docs/figures, scaled to fit the sheet width; returns the next free row."""
    p = os.path.join(FIG, fname)
    if not os.path.exists(p):
        return note(ws, row, f"(Figure {fname} not found — run make_docs.py first.)", n)
    w, h = PILImage.open(p).size
    avail = sum(width_of(ws, k) for k in range(1, n + 1)) * 7.0
    target = min(max_px, avail, w)
    img = XLImage(p)
    img.width, img.height = int(target), int(h * target / w)
    ws.add_image(img, f"A{row}")
    px, k = 0.0, row
    while px < img.height:          # rows under the image keep the default 15 pt (20 px) height
        ws.row_dimensions[k].height = 15
        px += 20
        k += 1
    return k + 2


def link(cell, url, text=None):
    cell.value = text or url
    cell.hyperlink = url
    cell.font = F(color="1C5CAB", underline="single")


# ============================================================================ Glossary
ws, r, n = new_sheet("Glossary", "Glossary: the terms used in this workbook",
                     "short, plain-language meanings of every technical term.",
                     "SeK is the main score. For every score in this workbook, higher is better "
                     "(except error and false-alarm rates).", [28, 72, 50])
terms = [
    ("Semantic change detection", "Compare two aerial photos of the same place taken at different times (T1 = before, "
     "T2 = after). For every pixel, decide whether it changed and, if so, what it was before and after "
     "(e.g. tree → building).", "The task this whole project solves."),
    ("SECOND dataset", "A public benchmark of 2,968 aerial image pairs with 6 land-cover classes: non-vegetated "
     "ground, tree, low vegetation, water, building, playground.", "Standard dataset, so results can be compared."),
    ("Class imbalance", "Some classes are much more common than others. Here only 20% of pixels change, and the "
     "largest class is 59× bigger than the smallest.", "Models tend to ignore rare classes."),
    ("OA (overall accuracy)", "Share of all pixels given the right label.",
     "Misleading here: predicting 'no change' everywhere already gets about 80%."),
    ("IoU", "Intersection over union: overlap between the predicted and the true area of a class, 0–100%.",
     "Standard per-class score."),
    ("mIoU", "Average of the IoU for 'changed' and for 'unchanged'.", "Measures how well change is located."),
    ("SeK (main score)", "Separated kappa: agreement on the changed classes, with the easy 'no change' part "
     "removed, scaled by how well change is found.", "Official main metric of SECOND; rewards both finding "
     "changes and naming the classes correctly."),
    ("Fscd", "Balance (harmonic mean) of precision and recall of correct classes on changed pixels.",
     "Second official metric."),
    ("Precision / recall", "Precision: of the pixels the model called 'changed', how many really changed. "
     "Recall: of the pixels that really changed, how many the model found.", "Shows the trade-off between "
     "false alarms and missed changes."),
    ("Missed change / false alarm", "Missed: a changed pixel predicted as 'no change'. False alarm: an unchanged "
     "pixel predicted as changed.", "The two main kinds of error."),
    ("Architecture", "The design of the neural network.", "Phase A compares three designs."),
    ("Early Fusion", "Stacks both images into one input for one network.", "Simplest design; the starting point."),
    ("Siamese (SSCD)", "The same encoder reads each image separately, then the features are compared.",
     "Keeps the two dates apart, which helps a lot."),
    ("Bi-SRNet-lite", "SSCD plus attention between the two dates and a consistency loss.", "Best design here."),
    ("ResNet-18", "A standard image network, pre-trained on ImageNet, used as the encoder in all three designs.",
     "The same encoder everywhere keeps the comparison fair."),
    ("Loss / CE", "The loss tells the model how wrong it is during training. CE (cross-entropy) is the "
     "standard one.", "Phase B changes the loss to handle imbalance."),
    ("Weighted CE, Median-freq, Class-balanced", "Versions of CE that give rare classes a bigger weight.",
     "Three ways to choose the weights."),
    ("Dice loss", "A loss based on overlap between prediction and truth; small classes count more.",
     "Part of the best combination."),
    ("Focal loss / OHEM", "Focus training on hard pixels (focal: smoothly; OHEM: only the hardest 25%).",
     "Did not help here."),
    ("Rare-class sampling", "Show training images that contain rare classes more often.", "A fix on the data side."),
    ("Epoch", "One pass over all training images. Every model here trains for 20 epochs.", "—"),
    ("Train / validation / test", "Train: used to learn. Validation: used to pick the best epoch. Test: used only "
     "once, for the final score (595 image pairs).", "Test scores are fair because the model never saw these "
     "images."),
    ("Cross-validation (3-fold)", "Split the training data into 3 parts; train 3 times, each time holding out a "
     "different part. Report mean ± standard deviation.", "Shows whether a result is stable or just luck."),
    ("Overfitting / underfitting", "Overfitting: much better on training data than on new data. Underfitting: "
     "not yet good even on training data.", "Checked with train-vs-validation curves."),
    ("Baseline", "The reference to compare against (Early Fusion + CE, or plain CE on the same design).", "—"),
]
_, _, r = table(ws, r, ["Term", "Meaning in plain words", "Why it matters here"], [list(t) for t in terms],
                wrap_cols=(1, 2))

# ============================================================================ Dataset
ws, r, n = new_sheet("Dataset", "Dataset: SECOND and its class imbalance",
                     "basic facts about the data and how unbalanced the classes are.",
                     f"Only {imb['change_pixel_ratio'] * 100:.0f}% of pixels change, and the largest class is "
                     f"{imb['semantic_imbalance_ratio_max_over_min']:.0f}× bigger than the smallest. Water and "
                     "playground are very rare, which makes them hard to learn.", [36, 16, 16, 20, 20, 14])
r = heading(ws, r, "Overview", n)
ov = [["Image pairs", imb["n_pairs"], "#,##0"], ["Image size used", "256 × 256 pixels (original 512 × 512)", None],
      ["Share of pixels that change", imb["change_pixel_ratio"], "0.0%"],
      ["Unchanged : changed pixels", imb["change_vs_nochange_ratio"], "0.00\" : 1\""],
      ["Largest class ÷ smallest class", imb["semantic_imbalance_ratio_max_over_min"], "0\"×\""],
      ["Types of change (before → after)", len(imb["transitions_pct"]), "0"],
      ["Train / validation / test pairs", "2,077 / 296 / 595 (random 70/10/20 split, seed 0)", None]]
f0, l0, r = table(ws, r, ["Item", "Value"], [o[:2] for o in ov])
for i, o in enumerate(ov):
    if o[2]:
        ws.cell(f0 + i, 2).number_format = o[2]
for k in range(f0 - 1, l0 + 1):  # the Value column spans B:E so long text fits on one line
    ws.merge_cells(start_row=k, start_column=2, end_row=k, end_column=5)
    ws.cell(k, 2).alignment = Alignment(horizontal="center" if k == f0 - 1 else "left", vertical="center")
r = heading(ws, r, "Share of changed pixels per class", n)
rows = [[c, imb["T1_class_pct_of_changed"][c] / 100, imb["T2_class_pct_of_changed"][c] / 100,
         imb["T1_images_containing_class"][c], imb["T2_images_containing_class"][c]] for c in SEM]
_, _, r = table(ws, r, ["Class", "Before (T1)", "After (T2)", "Images with class (T1)",
                        "Images with class (T2)"], rows, fmts=[None, "0.0%", "0.0%", "#,##0", "#,##0"])
r = picture(ws, r, "class_distribution.png", n, 800)
r = heading(ws, r, "Most common types of change (top 15 of 31)", n)
tr_rows = [[k.replace("->", "→"), v / 100] for k, v in list(imb["transitions_pct"].items())[:15]]
_, _, r = table(ws, r, ["Change (before → after)", "Share of changed pixels"], tr_rows, fmts=[None, "0.00%"])
r = note(ws, r, "The 5 most common changes cover over 70% of all changed pixels; the rarest (for example "
         "playground → water) are below 0.01%. The matrix below shows all 31 types.", n)
r = picture(ws, r, "transition_matrix.png", n, 620)

# ============================================================================ Architectures
ws, r, n = new_sheet("Architectures", "Phase A: which network design works best?",
                     "three network designs trained in exactly the same way (same encoder, data split, seed and "
                     "20 epochs), all with the standard CE loss.",
                     f"Bi-SRNet-lite is best (test SeK {base['t']['SeK']:.2f}). Changing the design alone raised "
                     f"SeK by {(base['t']['SeK'] - ef['t']['SeK']) / ef['t']['SeK']:.0%} over Early Fusion, the "
                     "biggest improvement in the whole study.",
                     [18, 30, 44, 11, 10, 10, 10, 10, 12, 13, 11])
hdr = ["Model", "Based on", "What it does", "Size (million params)", "SeK (%)", "Fscd (%)", "mIoU (%)", "OA (%)",
       "Train time (min)", "SeK gain vs Early Fusion", "Rank (by SeK)"]
rows = []
start = r + 1
for i, x in enumerate(A):
    rr = start + i
    m = ARCH[x["model"]]
    rows.append([m[0], m[1], m[3], m[2], x["t"]["SeK"], x["t"]["Fscd"], x["t"]["mIoU"], x["t"]["OA"], x["mins"],
                 f"=(E{rr}-$E${start})/$E${start}", f"=RANK(E{rr},$E${start}:$E${start + len(A) - 1})"])
f1, l1, r = table(ws, r, hdr, rows,
                  fmts=[None, None, None, "0.00", "0.00", "0.00", "0.00", "0.00", "0.0", "+0.0%;-0.0%;\"—\"", "0"],
                  best_row=[x["model"] for x in A].index(best_arch), wrap_cols=(1, 2))
REFS["arch"] = (f1, l1)
REFS["arch_best_row"] = f1 + [x["model"] for x in A].index(best_arch)
r = note(ws, r, "All scores are on the test split (595 image pairs the models never saw during training). "
         "Green row = design used for Phase B (chosen by validation score, not by test score). "
         "Gain and rank are formulas.", n)
r = picture(ws, r, "architecture_comparison.png", n, 1000)

# ============================================================================ Imbalance techniques
n_better = sum(x["t"]["SeK"] > base["t"]["SeK"] for x in B)
worse = [x["label"] for x in B if x["t"]["SeK"] <= base["t"]["SeK"]]
ws, r, n = new_sheet("Imbalance Techniques", "Phase B: which way of handling rare classes works best?",
                     f"nine techniques for class imbalance, all on {ARCH[best_arch][0]}, compared with the plain "
                     "CE loss (first row).",
                     f"{n_better} of 9 techniques beat plain CE. Best: {B[0]['label']} "
                     f"(+{(B[0]['t']['SeK'] - base['t']['SeK']) / base['t']['SeK']:.1%} SeK). "
                     f"{', '.join(worse)} did worse.",
                     [26, 40, 28, 9, 9, 9, 9, 11, 12, 12, 10, 11])
hdr = ["Technique", "What it does", "Paper", "SeK (%)", "Fscd (%)", "mIoU (%)", "OA (%)", "Water IoU (%)",
       "Playground IoU (%)", "SeK change vs CE", "Rank (by SeK)", "Better than CE?"]
rows = []
start = r + 1
for i, x in enumerate(BB):
    rr = start + i
    lab, desc, ref, _ = TECH[x["key"]]
    rows.append([lab, desc, ref, x["t"]["SeK"], x["t"]["Fscd"], x["t"]["mIoU"], x["t"]["OA"], x["pc"][4], x["pc"][6],
                 f"=(D{rr}-$D${start})/$D${start}", f"=RANK(D{rr},$D${start}:$D${start + len(BB) - 1})",
                 "— (baseline)" if i == 0 else f"=IF(D{rr}>$D${start},\"Yes\",\"No\")"])
f2, l2, r = table(ws, r, hdr, rows,
                  fmts=[None, None, None] + ["0.00"] * 6 + ["+0.0%;-0.0%;\"—\"", "0", None],
                  best_row=1, wrap_cols=(1, 2))
REFS["tech"] = (f2, l2)
for i, x in enumerate(BB):
    url = TECH[x["key"]][3]
    if url:
        c = ws.cell(f2 + i, 3)
        link(c, url, TECH[x["key"]][2])
        c.border = BOX
        c.alignment = Alignment(wrap_text=True, vertical="center")
for i in range(len(BB)):
    ws.cell(f2 + i, 12).alignment = Alignment(horizontal="center", vertical="center")
ws.conditional_formatting.add(f"J{f2 + 1}:J{l2}", ColorScaleRule(
    start_type="num", start_value=-0.08, start_color="F4C7C3", mid_type="num", mid_value=0,
    mid_color="FFFFFF", end_type="num", end_value=0.08, end_color="B7E1CD"))
r = note(ws, r, "Green row = best technique. 'SeK change vs CE' is relative to the first row (red = worse, "
         "green = better). OHEM did not train properly with its default setting (hardest 25% of pixels); this "
         "shows that setting failing, not that OHEM can never work. Each technique was trained once; differences "
         "below about 0.5 SeK should be treated as ties — see the Cross-Validation sheet.", n)
r = picture(ws, r, "imbalance_techniques.png", n, 800)
r = picture(ws, r, "relative_gain_vs_ce.png", n, 1000)

# ============================================================================ Per-class IoU
ws, r, n = new_sheet("Per-class IoU", "How well is each land-cover class recognised?",
                     "the IoU (overlap score, 0–100) of every class for every model and technique. Darker blue "
                     "means better.",
                     f"The rare classes improve the most: water goes from {ef['pc'][4]:.1f} (Early Fusion) to "
                     f"{best['pc'][4]:.1f}, and playground from {ef['pc'][6]:.1f} to {best['pc'][6]:.1f} "
                     "(best model).", [44] + [14] * 7)
rows = [[full(x)] + list(x["pc"]) for x in ALL]
f3, l3, r = table(ws, r, ["Model + technique"] + CLASSES, rows, fmts=[None] + ["0.0"] * 7)
ws.conditional_formatting.add(f"B{f3}:H{l3}", ColorScaleRule(
    start_type="num", start_value=0, start_color="FFFFFF", mid_type="num", mid_value=40,
    mid_color="9EC5F4", end_type="num", end_value=90, end_color="3987E5"))
r = note(ws, r, "First 3 rows: the three designs with plain CE. Other rows: imbalance techniques on "
         "Bi-SRNet-lite, sorted by SeK. 'no-change' is always easy (about 85) because it covers 80% of all pixels.", n)

# ============================================================================ Error analysis
miss_big = 0
for x in ALL:
    C = x["conf"].T
    miss_big += int(C[1:, 0].sum() > C[1:, 1:].sum() - np.trace(C[1:, 1:]))
ws, r, n = new_sheet("Error Analysis", "What kind of mistakes do the models make?",
                     "every truly changed pixel is either given the correct class, the wrong class, or missed "
                     "(the model says 'no change'). A false alarm is an unchanged pixel marked as changed.",
                     f"Missing a change is the biggest mistake in {miss_big} of {len(ALL)} runs. Imbalance "
                     "techniques find more changes (higher recall) but also raise more false alarms "
                     "(lower precision).", [40, 14, 13, 13, 13, 14, 13, 10, 11, 10, 11, 11, 11])
hdr = ["Model + technique", "Changed pixels (truth)", "Correct class", "Wrong class", "Missed change",
       "Unchanged pixels (truth)", "False alarms", "Correct %", "Wrong class %", "Missed %", "False alarm %",
       "Change precision", "Change recall"]
rows = []
start = r + 1
for i, x in enumerate(ALL):
    C = x["conf"].T.astype(np.int64)  # rows = truth, cols = prediction
    rr = start + i
    ok = int(np.trace(C[1:, 1:]))
    rows.append([full(x), int(C[1:].sum()), ok, int(C[1:, 1:].sum() - ok), int(C[1:, 0].sum()), int(C[0].sum()),
                 int(C[0, 1:].sum()), f"=C{rr}/B{rr}", f"=D{rr}/B{rr}", f"=E{rr}/B{rr}", f"=G{rr}/F{rr}",
                 f"=(C{rr}+D{rr})/(C{rr}+D{rr}+G{rr})", f"=(C{rr}+D{rr})/B{rr}"])
f4, l4, r = table(ws, r, hdr, rows, fmts=[None] + ["#,##0"] * 6 + ["0.0%"] * 6)
ws.conditional_formatting.add(f"J{f4}:J{l4}", ColorScaleRule(
    start_type="min", start_color="FFFFFF", end_type="max", end_color="F4B183"))
r = note(ws, r, "Pixel counts come from each run's test confusion matrix (both dates together). The percentage "
         "columns are formulas. Orange shading = share of changes that were missed.", n)
r = picture(ws, r, "error_breakdown.png", n, 1000)
r = picture(ws, r, "confusion_best.png", n, 600)

# ============================================================================ Cross-validation
cv_note_best = ""
if len(cv_complete) >= 2:
    b0 = max(cv_complete, key=lambda t: np.mean(CVS[t]["test"]))
    cv_note_best = f"Highest mean so far: {cv_label(b0)} ({np.mean(CVS[b0]['test']):.2f} ± " \
                   f"{np.std(CVS[b0]['test'], ddof=1):.2f})."
wins_b = [t for t in cv_complete if t.startswith("B_") and "A_bisrnet_ce" in CVS
          and sum(a > b for a, b in zip(CVS[t]["test"], CVS["A_bisrnet_ce"]["test"])) == K]
wins_txt = (", ".join(cv_label(t).split(" + ", 1)[1] for t in wins_b) + " beat plain CE in all 3 folds. ") \
    if wins_b else ""
ws, r, n = new_sheet("Cross-Validation", "Are the results stable? (3-fold cross-validation)",
                     "each configuration is trained 3 times on different parts of the training data "
                     "(1,582 train / 791 validation pairs each time) and scored on the same test split. "
                     f"Finished: {cv_done} of {K * len(CV_ORDER)} runs.",
                     "The order of the three designs holds in every fold. " + wins_txt + cv_note_best
                     + " Techniques within about 0.3 SeK of each other are effectively tied.",
                     [40, 9, 9, 9, 9, 11, 9, 11, 11, 11, 11, 16, 44])
hdr = ["Configuration", "Fold 1", "Fold 2", "Fold 3", "Folds done", "Mean test SeK", "Std", "Mean val SeK",
       "Mean train SeK", "Gap (train − val)", "Val-loss rise", "Beats reference in", "Fit diagnosis"]
rows = []
start = r + 1
rowof = {t: start + i for i, t in enumerate(CV_ORDER)}
for t in CV_ORDER:
    rr = rowof[t]
    folds = CV.get(t, {})
    tests = [folds[k]["test"]["SeK"] if k in folds else None for k in range(K)]
    s = CVS.get(t)
    ref = "A_early_fusion_ce" if t.startswith("A_") else "A_bisrnet_ce"
    q = rowof[ref]
    win = "— (reference)" if ref == t else \
        f"=IF(AND(E{rr}=3,E{q}=3),((B{rr}>B{q})+(C{rr}>C{q})+(D{rr}>D{q}))&\" of 3 folds\",\"—\")"
    diag = (s["diag"] + ("; best epoch was the last, so more epochs may help" if s["last_best"] else "")) if s \
        else ("running" if folds else "not started yet")
    rows.append([cv_label(t)] + tests + [
        f"=COUNT(B{rr}:D{rr})", f"=IF(E{rr}>0,AVERAGE(B{rr}:D{rr}),\"—\")", f"=IF(E{rr}>1,STDEV(B{rr}:D{rr}),\"—\")",
        s["val"] if s else None, s["train"] if s else None,
        f"=IF(AND(ISNUMBER(I{rr}),ISNUMBER(H{rr})),I{rr}-H{rr},\"—\")", s["rise"] if s else None, win, diag])
f5, l5, r = table(ws, r, hdr, rows,
                  fmts=[None, "0.00", "0.00", "0.00", "0", "0.00", "0.00", "0.00", "0.00", "0.00", "0.0%", None, None],
                  wrap_cols=(12,))
for i in range(len(CV_ORDER)):
    ws.cell(f5 + i, 12).alignment = Alignment(horizontal="center", vertical="center")
r = note(ws, r, "Reference: the three designs are compared with Early Fusion; the techniques with "
         "Bi-SRNet-lite + CE. 'Beats reference' counts the folds where the configuration scored higher. "
         "Fit diagnosis rules (averaged over folds): Overfitting = validation loss ends more than 10% above its "
         "lowest value AND train SeK is more than 3 points above validation SeK; Mild overfitting = only one of "
         "the two; Underfitting = validation SeK still rising by more than 0.15 per epoch with a gap below 2; "
         "Failed = best validation SeK below 5. These thresholds were chosen for this study. CV models train on "
         "fewer images (1,582) than the main runs (2,077), so their scores are slightly lower; CV is used for "
         "stability and ranking.", n)
r = picture(ws, r, "cv_scores.png", n, 900)
r = picture(ws, r, "generalization_gap.png", n, 900)
r = heading(ws, r, "Every fold in detail", n)
det = []
for t in CV_ORDER:
    for k, rj in sorted(CV.get(t, {}).items()):
        h = rj["history"][rj["best_epoch"] - 1]
        det.append([cv_label(t), k + 1, rj["best_epoch"], h.get("train_SeK"), h["SeK"], rj["test"]["SeK"],
                    rj["test"]["Fscd"], rj["test"]["mIoU"], rj["test"]["OA"]])
_, _, r = table(ws, r, ["Configuration", "Fold", "Best epoch", "Train SeK", "Val SeK", "Test SeK", "Test Fscd",
                        "Test mIoU", "Test OA"], det, fmts=[None, "0", "0"] + ["0.00"] * 6)

# ============================================================================ Learning curves
diag_counts = {}
for t in cv_complete:
    diag_counts[CVS[t]["diag"]] = diag_counts.get(CVS[t]["diag"], 0) + 1
ws, r, n = new_sheet("Learning Curves", "Overfitting or underfitting? (train vs validation curves)",
                     "the score on training data (orange) and on unseen validation data (blue) after every epoch, "
                     "averaged over the 3 folds. A growing gap between the two lines means overfitting; both lines "
                     "low or still climbing means underfitting.",
                     ("So far: " + ", ".join(f"{v} × {k.lower()}" for k, v in diag_counts.items()) + ". The "
                      "validation score flattens after about 12–15 epochs while the training score keeps rising, "
                      "so longer training alone is unlikely to help much.")
                     if cv_complete else "Cross-validation still running.",
                     [8] + [12] * 11)
r = picture(ws, r, "learning_curves_sek.png", n, 1000)
r = picture(ws, r, "learning_curves_loss.png", n, 1000)
if cv_complete:
    r = heading(ws, r, "Curve data (mean of 3 folds)", n)
    E = len(CVS[cv_complete[0]]["rs"][0]["history"])
    hdr = ["Epoch"]
    cols = []
    for t in cv_complete:
        short = cv_label(t).replace("Bi-SRNet-lite", "BiSR").replace(" (baseline)", "")
        for key, nm in (("train_SeK", "train SeK"), ("SeK", "val SeK")):
            hdr.append(f"{short}: {nm}")
            cols.append(np.mean([[h[key] for h in r_["history"]] for r_ in CVS[t]["rs"]], 0))
    for j in range(len(hdr)):
        ws.column_dimensions[L(j + 1)].width = 8 if j == 0 else 12
    rows = [[e + 1] + [float(c[e]) for c in cols] for e in range(E)]
    _, _, r = table(ws, r, hdr, rows, fmts=["0"] + ["0.00"] * len(cols))

# ============================================================================ Training history
ws, r, n = new_sheet("Training History", "Validation score during training (main runs)",
                     "the validation SeK after every epoch for the 12 main runs. The epoch with the best "
                     "validation SeK was kept and scored once on the test split.",
                     "Siamese designs learn faster and reach higher scores than Early Fusion; all runs level off "
                     "in the last few epochs.", [8] + [12] * 12)
r = picture(ws, r, "training_curves.png", n, 800)
E = len(ALL[0]["hist"])
hdr = ["Epoch"] + [full(x).replace("Bi-SRNet-lite", "BiSR") for x in ALL]
rows = [[e + 1] + [x["hist"][e]["SeK"] for x in ALL] for e in range(E)]
_, _, r = table(ws, r, hdr, rows, fmts=["0"] + ["0.00"] * len(ALL))

# ============================================================================ Published comparison
reach = {"early_fusion": 10.69, "sscd": 21.86, "bisrnet": 23.22}
ws, r, n = new_sheet("Published Comparison", "How do our scores compare with published papers?",
                     "our results next to results reported in papers on SECOND. The papers use 512 × 512 images, "
                     "bigger networks and longer training, so this is a reference, not a like-for-like contest.",
                     f"Our models reach {min(x['t']['SeK'] / reach[x['model']] for x in A):.0%}–"
                     f"{max(x['t']['SeK'] / reach[x['model']] for x in A):.0%} of the published score of the "
                     "same design, using ¼ of the pixels, a small ResNet-18 and 20 epochs on a laptop. The order "
                     "of the designs matches the papers.", [42, 34, 10, 10, 10, 10, 16])
rows = [[m, "Published (512 px)", oa, miou, sek, fs, "—"] for m, oa, miou, sek, fs in LIT]
start = r + 1
for x in A + [best]:
    li = {"early_fusion": 0, "sscd": 2, "bisrnet": 3}[x["model"]]
    rr = start + len(rows)
    rows.append([f"Ours: {full(x)}", "Ours (256 px, ResNet-18, 20 epochs)", x["t"]["OA"], x["t"]["mIoU"],
                 x["t"]["SeK"], x["t"]["Fscd"], f"=E{rr}/E{start + li}"])
f7, l7, r = table(ws, r, ["Method", "Setting", "OA (%)", "mIoU (%)", "SeK (%)", "Fscd (%)",
                          "Share of published SeK"], rows, fmts=[None, None, "0.00", "0.00", "0.00", "0.00", "0%"])
for i in range(len(rows)):
    ws.cell(f7 + i, 7).alignment = Alignment(horizontal="center", vertical="center")
r = note(ws, r, "Published values as tabulated in Mamba-FCS (arXiv:2508.08232), Table II. 'Share of published "
         "SeK' compares each of our models with the paper version of the same design (Early Fusion ↔ "
         "HRSCD-str.2, SSCD ↔ SSCD-l, Bi-SRNet-lite ↔ Bi-SRNet).", n)
r = picture(ws, r, "literature_comparison.png", n, 800)

# ============================================================================ Predictions
ws, r, n = new_sheet("Predictions", "What the best model actually predicts",
                     "four test image pairs: the two photos, the true labels, and the best model's prediction. "
                     "White = no change; colours = land-cover class (legend under the picture).",
                     "Large changes such as new buildings are found well; thin or small changed areas are often "
                     "missed, which matches the error analysis.", [12] * 12)
r = picture(ws, r, "qualitative.png", n, 1000)

# ============================================================================ Methods & references
ws, r, n = new_sheet("Methods & References", "Every model and technique used, with its source",
                     "the paper behind each component, with a clickable link. Full details and code locations: "
                     "docs/REFERENCES.md in the repository.",
                     "All models and techniques come from published work; this project's contribution is the "
                     "controlled comparison and the error analysis.", [18, 46, 36, 54])
refs = [
    ("Dataset & metric", "SECOND dataset; SeK, Fscd, mIoU, OA", "Yang et al., IEEE TGRS 2022", "https://arxiv.org/abs/2010.05687"),
    ("Architecture", "Early Fusion (HRSCD strategy 2)", "Daudt et al., CVIU 2019", "https://arxiv.org/abs/1810.08452"),
    ("Architecture", "SSCD (SSCD-l)", "Ding et al., IEEE TGRS 2022", "https://arxiv.org/abs/2108.06103"),
    ("Architecture", "Bi-SRNet-lite (Bi-SRNet) and consistency loss", "Ding et al., IEEE TGRS 2022", "https://arxiv.org/abs/2108.06103"),
    ("Encoder", "ResNet-18, ImageNet pre-trained", "He et al., CVPR 2016", "https://arxiv.org/abs/1512.03385"),
    ("Decoder", "Feature Pyramid Network style", "Lin et al., CVPR 2017", "https://arxiv.org/abs/1612.03144"),
    ("Attention", "Attention between the two dates (non-local style)", "Wang et al., CVPR 2018", "https://arxiv.org/abs/1711.07971"),
    ("Attention gate", "Zero-initialised residual gate", "Zhang et al. (SAGAN), ICML 2019", "https://arxiv.org/abs/1805.08318"),
    ("Imbalance", "Median-frequency balancing", "Eigen & Fergus, ICCV 2015", "https://arxiv.org/abs/1411.4734"),
    ("Imbalance", "Class-balanced loss (effective number)", "Cui et al., CVPR 2019", "https://arxiv.org/abs/1901.05555"),
    ("Imbalance", "Focal loss", "Lin et al., ICCV 2017", "https://arxiv.org/abs/1708.02002"),
    ("Imbalance", "Dice loss", "Milletari et al. (V-Net), 3DV 2016", "https://arxiv.org/abs/1606.04797"),
    ("Imbalance", "OHEM (hard-example mining)", "Shrivastava et al., CVPR 2016", "https://arxiv.org/abs/1604.03540"),
    ("Imbalance", "Rare-class oversampling (repeat-factor style)", "Gupta et al. (LVIS), CVPR 2019", "https://arxiv.org/abs/1908.03195"),
    ("Imbalance", "Weighted CE (1/√frequency)", "Common practice", "—"),
    ("Training", "AdamW optimiser", "Loshchilov & Hutter, ICLR 2019", "https://arxiv.org/abs/1711.05101"),
    ("Training", "One-cycle learning-rate schedule", "Smith & Topin, 2017", "https://arxiv.org/abs/1708.07120"),
    ("Validation", "K-fold cross-validation", "Kohavi, IJCAI 1995", "https://www.ijcai.org/Proceedings/95-2/Papers/016.pdf"),
    ("Comparison only", "SCanNet", "Ding et al., IEEE TGRS 2024", "https://arxiv.org/abs/2212.05245"),
    ("Comparison only", "ChangeMamba", "Chen et al., IEEE TGRS 2024", "https://arxiv.org/abs/2404.03425"),
    ("Comparison only", "Mamba-FCS (published table; SeK loss)", "arXiv 2025", "https://arxiv.org/abs/2508.08232"),
    ("Comparison only", "FD-ProtoSCD (imbalanced change types)", "Remote Sensing 2026", "https://www.mdpi.com/2072-4292/18/17/2957"),
]
f6, _, r = table(ws, r, ["Type", "Component", "Paper", "Link"], [[a, b, c, d] for a, b, c, d in refs],
                 wrap_cols=(1, 2))
for i, (_, _, _, u) in enumerate(refs):
    if u.startswith("http"):
        c = ws.cell(f6 + i, 4)
        link(c, u)
        c.border = BOX
        c.alignment = Alignment(vertical="center")
r = note(ws, r, "Details for the dataset and change-detection papers were checked against their arXiv pages. "
         "Confirm the venue and year of the others before citing them in a paper.", n)

# ============================================================================ Setup
ws, r, n = new_sheet("Setup", "How the experiments were run",
                     "the settings shared by every run, and exact definitions of the scores.",
                     "Every run used the same data split, seed, encoder and training schedule, so differences "
                     "between rows come only from the design or technique being tested.", [24, 110])
setup = [
    ("Data", "SECOND public training release: 2,968 image pairs, resized from 512 to 256 pixels"),
    ("Split", "Random 70/10/20 (seed 0): 2,077 train / 296 validation / 595 test"),
    ("Cross-validation", "Train + validation (2,373 pairs) re-split into 3 folds (seed 1); test split unchanged"),
    ("Augmentation", "Random horizontal/vertical flips, 90° rotations, and swapping the before/after images"),
    ("Encoder", "ResNet-18 pre-trained on ImageNet (torchvision)"),
    ("Optimiser", "AdamW, weight decay 0.0001; learning rate 0.0005 (decoders), 0.0001 (encoder)"),
    ("Schedule", "One-cycle schedule with 10% warm-up; 20 epochs; batch size 8"),
    ("Baseline loss", "BCE (change) + 0.5 × CE (class, only on changed pixels) [+ consistency loss for Bi-SRNet-lite]"),
    ("Model selection", "Keep the epoch with the best validation SeK; score it once on the test split"),
    ("Hardware", "Apple M4 laptop, 16 GB RAM, PyTorch on the Apple GPU (MPS); about 10 min (Early Fusion) to 22 min per run"),
    ("Software", "Python 3.11, PyTorch 2.14, torchvision 0.29 (see requirements.txt)"),
    ("OA", "Share of all pixels with the correct label (7 labels: no-change + 6 classes)"),
    ("mIoU", "Mean of the IoU of 'no change' and the IoU of 'change'"),
    ("SeK (main score)", "Kappa over the changed classes (the no-change/no-change cell removed) × e^(IoU_change − 1)"),
    ("Fscd", "Harmonic mean of precision and recall of correct classes on changed pixels"),
    ("Code", "github.com/Chrishabh2002/Second (train.py, scd.py, run_all.sh, run_cv.sh)"),
]
_, _, r = table(ws, r, ["Item", "Value"], [list(s) for s in setup], wrap_cols=(1,))

# ============================================================================ Next steps
ws, r, n = new_sheet("Next Steps", "Next steps towards a publication",
                     "the plan, in order of priority. Full plan: docs/ROADMAP.md in the repository.",
                     "The main missing piece is a new method. The planned one is a transition-aware loss that "
                     "targets the biggest error found here: missed changes on rare types of change.",
                     [6, 62, 54, 22])
steps = [
    (1, "Finish cross-validation; add 2 more seeds for the best configurations", "Every claim gets a mean ± std", "About 1 day"),
    (2, "Tune the change threshold on the validation split", "Missed changes are the largest error", "Hours, no retraining"),
    (3, "Train the best configuration for 40 epochs", "In CV its best epoch was often the last one", "About 45 min"),
    (4, "Fix or drop OHEM (keep 50–70% of pixels)", "The default setting failed", "2–3 runs"),
    (5, "Transition-aware loss (new method)", "Weights each pixel by how rare its before → after change is", "1–2 weeks"),
    (6, "Published setting on a GPU (512 px, larger encoder, 50+ epochs)", "Needed to compare with published models", "About 1 week"),
    (7, "Second dataset (Landsat-SCD)", "Shows the gain is not specific to SECOND", "About 1 week"),
    (8, "Write the paper (IEEE GRSL / IGARSS / MDPI Remote Sensing)", "Venue to be chosen with the supervisor", "2–3 weeks"),
]
_, _, r = table(ws, r, ["#", "Step", "Why", "Time"], [list(s) for s in steps], wrap_cols=(1, 2))
for i in range(len(steps)):
    ws.cell(r - 2 - len(steps) + i, 1).alignment = Alignment(horizontal="center", vertical="center")

# ============================================================================ Summary (built last, shown first)
a0, _ = REFS["arch"]
t0, t1 = REFS["tech"]
ws, r, n = new_sheet("Summary", "Semantic change detection on SECOND — results summary",
                     "the headline numbers and findings. Each other sheet explains one part in detail (links at "
                     "the bottom). All scores are on a held-out test set of 595 image pairs; higher is better.",
                     f"The best model ({ARCH[best['model']][0]} + {best['label']}) reaches test SeK "
                     f"{best['t']['SeK']:.2f}, {(best['t']['SeK'] - ef['t']['SeK']) / ef['t']['SeK']:.0%} higher "
                     f"than the starting model ({ef['t']['SeK']:.2f}). Most of the gain comes from the network "
                     "design; handling class imbalance adds a smaller extra gain.", [54, 20, 72])
r = heading(ws, r, "Key numbers", n)
k0 = r + 1
kpis = [
    ("Starting model: Early Fusion + CE — test SeK (%)", f"=Architectures!E{a0}", "0.00", "The simplest design"),
    ("Best design: Bi-SRNet-lite + CE — test SeK (%)", f"=Architectures!E{REFS['arch_best_row']}", "0.00",
     "Best of the 3 designs"),
    ("Best overall — test SeK (%)", f"=MAX('Imbalance Techniques'!D{t0}:D{t1})", "0.00", "Best of all 12 runs"),
    ("Best technique", f"=INDEX('Imbalance Techniques'!A{t0}:A{t1},MATCH(B{k0 + 2},'Imbalance Techniques'!D{t0}:D{t1},0))",
     None, "Used on Bi-SRNet-lite"),
    ("Total improvement over the starting model", f"=(B{k0 + 2}-B{k0})/B{k0}", "0.0%", "Design + technique"),
    ("Improvement from the design alone", f"=(B{k0 + 1}-B{k0})/B{k0}", "0.0%", "Early Fusion → Bi-SRNet-lite"),
    ("Extra improvement from the best technique", f"=(B{k0 + 2}-B{k0 + 1})/B{k0 + 1}", "0.0%",
     "Same design: plain CE → best technique"),
    ("Techniques that beat plain CE", f"=COUNTIF('Imbalance Techniques'!L{t0 + 1}:L{t1},\"Yes\")&\" of 9\"", None,
     "On test SeK"),
    ("Cross-validation runs finished", f"=\"{cv_done} of {K * len(CV_ORDER)}\"", None, "3 folds × 12 configurations"),
]
_, _, r = table(ws, r, ["Measure", "Value", "Meaning"], [[a, b, d] for a, b, _, d in kpis], wrap_cols=(2,))
for i, (_, _, fm, _) in enumerate(kpis):
    c = ws.cell(k0 + i, 2)
    c.font = F(size=12, bold=True, color=NAVY)
    c.alignment = Alignment(horizontal="right", vertical="center")
    if fm:
        c.number_format = fm
    ws.row_dimensions[k0 + i].height = 20
r = picture(ws, r, "sek_journey.png", n, 800)
r = heading(ws, r, "Main findings", n)
cv_line = ("Cross-validation (3 folds) confirms the order of the three designs in every fold. " + wins_txt
           + cv_note_best + " Techniques this close should be treated as ties.") \
    if len(cv_complete) >= 2 else "Cross-validation is still running."
findings = [
    (1, "The network design matters most. Moving from Early Fusion to a Siamese design, then adding attention "
          f"between the two dates (Bi-SRNet-lite), raises SeK from {ef['t']['SeK']:.2f} to {base['t']['SeK']:.2f}."),
    (2, f"Handling class imbalance adds a smaller gain. {n_better} of 9 techniques beat plain CE; the best is "
          f"{B[0]['label']} ({B[0]['t']['SeK']:.2f}). Focal loss and rare-class sampling alone did not help; OHEM "
          "failed with its default setting."),
    (3, f"Rare classes gain the most: water IoU {ef['pc'][4]:.1f} → {best['pc'][4]:.1f}, playground "
          f"{ef['pc'][6]:.1f} → {best['pc'][6]:.1f}."),
    (4, "The biggest remaining error is missed changes (changed pixels predicted as 'no change'), not mixing up "
          "classes. Imbalance techniques find more changes but also raise false alarms."),
    (5, cv_line),
    (6, "Scores are below published papers (e.g. Bi-SRNet 23.22 SeK) because this study uses 256-pixel images, "
          "a small ResNet-18 and 20 epochs on a laptop. The comparison inside this study is fair; the one with "
          "papers is only a reference."),
]
fr0 = r + 1
_, _, r = table(ws, r, ["#", "Finding"], [list(f_) for f_ in findings], wrap_cols=(1,))
ws.merge_cells(start_row=fr0 - 1, start_column=2, end_row=fr0 - 1, end_column=3)
for i, (_, txt) in enumerate(findings):  # the finding text spans columns B:C
    k = fr0 + i
    ws.merge_cells(start_row=k, start_column=2, end_row=k, end_column=3)
    ws.cell(k, 1).alignment = Alignment(horizontal="center", vertical="top")
    ws.cell(k, 2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[k].height = max(18, lines_needed(txt, (20 + 72) * 1.05 - 4) * PT_PER_LINE + 6)
r = heading(ws, r, "Sheets in this workbook (click a name to open it)", n)
contents = [
    ("Glossary", "Plain-language meaning of every term (SeK, IoU, Dice, cross-validation, …)"),
    ("Dataset", "The data and how unbalanced its classes are"),
    ("Architectures", "Phase A: 3 network designs compared"),
    ("Imbalance Techniques", "Phase B: 9 ways of handling rare classes"),
    ("Per-class IoU", "How well each land-cover class is recognised"),
    ("Error Analysis", "What kind of mistakes the models make"),
    ("Cross-Validation", "Are the results stable? (3-fold cross-validation)"),
    ("Learning Curves", "Overfitting or underfitting?"),
    ("Training History", "Validation score after every epoch"),
    ("Published Comparison", "Our scores next to published papers"),
    ("Predictions", "Example outputs of the best model"),
    ("Methods & References", "Every component with its paper and link"),
    ("Setup", "Settings and exact metric definitions"),
    ("Next Steps", "Plan towards a publication"),
]
c0 = r + 1
_, _, r = table(ws, r, ["Sheet", "What you will find"], [[a, b] for a, b in contents], wrap_cols=(1,))
ws.merge_cells(start_row=c0 - 1, start_column=2, end_row=c0 - 1, end_column=3)
for i, (name, _) in enumerate(contents):
    c = ws.cell(c0 + i, 1)
    link(c, f"#'{name}'!A1", name)
    c.border = BOX
    ws.merge_cells(start_row=c0 + i, start_column=2, end_row=c0 + i, end_column=3)
    ws.row_dimensions[c0 + i].height = 18
r = note(ws, r, "Source: results/*.json written by train.py (code: github.com/Chrishabh2002/Second). Scores use "
         "the official SECOND metrics; SeK is the main one. Measured values come straight from the result files; "
         "gains, ranks, means and counts are Excel formulas.", n)

# ----------------------------------------------------------------------------- order, save, cache values
order = ["Summary", "Glossary", "Dataset", "Architectures", "Imbalance Techniques", "Per-class IoU",
         "Error Analysis", "Cross-Validation", "Learning Curves", "Training History", "Published Comparison",
         "Predictions", "Methods & References", "Setup", "Next Steps"]
wb._sheets = [wb[s] for s in order]
for w in wb.worksheets:
    w.sheet_properties.tabColor = {"Summary": "0D366B", "Glossary": "1BAF7A",
                                   "Next Steps": "EB6834"}.get(w.title, "6DA7EC")
wb.active = 0
wb.calculation.fullCalcOnLoad = True
if os.environ.get("PREVIEW_SHEET"):  # used only to render a preview of one sheet
    wb.move_sheet(os.environ["PREVIEW_SHEET"], offset=-wb.sheetnames.index(os.environ["PREVIEW_SHEET"]))
    wb.active = 0
    OUT = os.environ["PREVIEW_OUT"]
wb.save(OUT)


def cache_values(path):
    """Store each formula's computed value in the file, so previews that do not recalculate
    (Quick Look, mail and Drive previews) show numbers instead of 0. Excel still recalculates."""
    import re
    import shutil
    import tempfile
    import zipfile
    from xml.sax.saxutils import escape

    import formulas
    sol = formulas.ExcelModel().loads(path).finish().calculate()
    fname = os.path.basename(path)
    zin = zipfile.ZipFile(path)
    wbx = zin.read("xl/workbook.xml").decode()
    rels = zin.read("xl/_rels/workbook.xml.rels").decode()
    target = {m.group(1): m.group(2) for m in re.finditer(r'Id="(rId\d+)"[^>]*Target="/?(?:xl/)?([^"]+)"', rels)}
    target.update({m.group(2): m.group(1) for m in re.finditer(r'Target="/?(?:xl/)?([^"]+)"[^>]*Id="(rId\d+)"', rels)})
    sheets = {"xl/" + target[m.group(2)]: m.group(1)
              for m in re.finditer(r'<sheet name="([^"]+)"[^>]*r:id="(rId\d+)"', wbx)}
    n_ok, bad = 0, []

    def fill(xml, title):
        def rep(m):
            nonlocal n_ok
            attrs, ref, f = m.group(1), m.group(2), m.group(3)
            v = sol.get(f"'[{fname}]{title.upper()}'!{ref}")
            v = v.value[0, 0] if v is not None and hasattr(v, "value") else v
            if v is None or isinstance(v, formulas.tokens.operand.XlError):
                bad.append(f"{title}!{ref}")
                return m.group(0)
            n_ok += 1
            if isinstance(v, str):
                return f'<c{attrs} t="str"><f>{f}</f><v>{escape(v)}</v></c>'
            if isinstance(v, (bool, np.bool_)):
                return f'<c{attrs} t="b"><f>{f}</f><v>{int(v)}</v></c>'
            return f"<c{attrs}><f>{f}</f><v>{float(v)!r}</v></c>"
        return re.sub(r'<c( r="([A-Z]+\d+)"[^>]*?)><f>(.*?)</f><v\s*/></c>', rep, xml)

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx").name
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in sheets:
                data = fill(data.decode(), sheets[item.filename]).encode()
            zout.writestr(item, data)
    zin.close()
    shutil.move(tmp, path)
    if bad:
        raise SystemExit(f"formula errors in: {bad[:20]}")
    return n_ok


print(f"cached {cache_values(OUT)} formula values")
print(f"wrote {OUT}  ({len(wb.worksheets)} sheets, CV runs {cv_done}/{K * len(CV_ORDER)})")
